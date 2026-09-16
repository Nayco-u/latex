"""Genera content/05-anexos-tablas.tex a partir de 07-TraceabilityRequirimentsMatrix.xlsx.

Vuelve a correrlo si el Excel cambia:  python3 content/gen_tablas.py
"""
import re
import openpyxl

import os

def _find(nombre):
    """Busca el Excel en la carpeta padre del proyecto (PRACTICA2/)."""
    for base in ("..", ".", os.path.dirname(__file__), "/mnt/user-data/uploads/PRACTICA2"):
        ruta = os.path.join(base, nombre)
        if os.path.exists(ruta):
            return ruta
    raise FileNotFoundError(nombre + " no encontrado; ajusta la ruta a mano.")


XLSX = _find("07-TraceabilityRequirimentsMatrix.xlsx")
OUT = os.path.join(os.path.dirname(__file__), "05-anexos-tablas.tex")

# =====================================================================
#  CORRECCIÓN 1: mitigaciones desalineadas en la hoja RISK del Excel
# =====================================================================
#  En la hoja RISK, la columna "Mitigación Propuesta" quedó corrida respecto
#  de la columna ID desde SR-08 hacia abajo. Tal como está el Excel, SR-13
#  (consumo de 10 W) aparece con "Pruebas de compatibilidad CAN con OBC real"
#  y SR-15 (masa y centro de gravedad) con "Presupuesto de potencia", que no
#  corresponden. El diccionario siguiente reasigna cada mitigación al
#  requerimiento que efectivamente mitiga.
#
#  ESTO ES UN PARCHE. Lo correcto es arreglar el Excel (hoja RISK, columna
#  "Mitigación Propuesta" del bloque de SR, filas 9 a 17) y luego borrar este
#  diccionario.
MIT_OVERRIDE = {
    "SR-08": "Pruebas de compatibilidad CAN con OBC real",
    "SR-09": "Validación de lógica de estados y temporización",
    "SR-13": "Presupuesto de potencia con margen ≥25%",
    "SR-14": "Revisión de layout 3D y mockup mecánico",
    "SR-15": "Simulaciones CAD para estimaciones de masas y revisión periódica "
             "de masa en el prototipo",
    "SR-16": "Mediciones EMI/EMC y separación RF",
    "SR-17": "",
    "SR-18": "",
}

# =====================================================================
#  CORRECCIÓN 2: erratas de redacción heredadas del Excel
# =====================================================================
#  El manual ING2001 objeta el informe completo a partir de 10 faltas de
#  ortografía, y los anexos cuentan. Estas sustituciones corrigen la
#  redacción sin alterar el contenido técnico del requerimiento.
#  Conviene aplicarlas también en el Excel para no arrastrarlas de nuevo.
ERRATAS = [
    ("durante su orbita", "durante su órbita"),
    ("funcionamiento en órbita probada", "funcionamiento en órbita probado"),
    ("con una tamaño desde", "con un tamaño desde"),
    ("estaciones remotas ,", "estaciones remotas,"),
    ("1.000.000 mensajes de durante", "1.000.000 de mensajes durante"),
    ("eje logitudinal", "eje longitudinal"),
    ("despliege", "despliegue"),
    ("revisión periodica", "revisión periódica"),
    ("prueba térmo-vacio", "pruebas termo-vacío"),
    ("El sistema deberá ser respetar la envolvente interna permita, "
     "los volúmen prohibídos",
     "El sistema deberá respetar la envolvente interna permitida, "
     "los volúmenes prohibidos"),
    ("que uitlice el sistema", "que utilice el sistema"),
    # Separador decimal: el informe está en español y usa coma.
    ("0.01%", "0,01%"),
    ("3.0 kg", "3,0 kg"),
    ("1.0%", "1,0%"),
    ("0.1%", "0,1%"),
]

REPL = [
    ("\\", r"\textbackslash{}"), ("&", r"\&"), ("%", r"\%"), ("$", r"\$"),
    ("#", r"\#"), ("_", r"\_"), ("{", r"\{"), ("}", r"\}"),
    ("~", r"\textasciitilde{}"), ("^", r"\textasciicircum{}"),
    ("°", r"$^{\circ}$"), ("±", r"$\pm$"), ("≥", r"$\geq$"), ("≤", r"$\leq$"),
    ("<", r"$<$"), (">", r"$>$"), ("×", r"$\times$"),
    ("“", "``"), ("”", "''"), ('"', "''"), ("’", "'"),
]

# Superíndices Unicode del Excel (p. ej. "10⁻⁵ Torr"). Se convierten a modo
# matemático DESPUÉS del escapado, porque REPL escaparía el acento circunflejo.
_SUP_MAP = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻", "0123456789+-")
_SUP_RX = re.compile(r"[⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻]+")


def superindices(s):
    return _SUP_RX.sub(lambda m: "$^{%s}$" % m.group(0).translate(_SUP_MAP), s)


# Palabras en otro idioma: el informe las exige en cursiva. Se aplica DESPUÉS
# del escapado LaTeX, para no romper las llaves de \textit{}.
EXTRANJERISMOS = [
    r"Store \\& Forward", "Watchdog Timers", "Single Event Latch-up",
    "Single Event Upset", "Inrush Current", "Stress screening", "Power Bus Noise",
    "Eigenfrequencies", "Random Vib", "Sine Sweep", "Burn-wire", "Bake-out",
    "Outgassing", "Failsafe", "Watchdog", "Latch-up", "checksum", "firmware",
    "hardware", "software", "shaker", "downlink", "uplink", "layout", "mockup",
    "deployer", "Transponder", "Shock",
]
EXTRANJERISMOS.sort(key=len, reverse=True)
_ITAL_RX = [(t, re.compile(r"(?<![\w\\])" + t.replace(" ", r"\s+") + r"(?![\w])"))
            for t in EXTRANJERISMOS]


def cursivas(s):
    hechos = []

    def guardar(m):
        hechos.append(m.group(0))
        return "\x00%d\x00" % (len(hechos) - 1)

    s = re.sub(r"\\textit\{[^{}]*\}", guardar, s)

    def envolver(m):
        hechos.append(r"\textit{%s}" % m.group(0))
        return "\x00%d\x00" % (len(hechos) - 1)

    for _, rx in _ITAL_RX:
        s = rx.sub(envolver, s)
    for i, v in enumerate(hechos):
        s = s.replace("\x00%d\x00" % i, v)
    return s


def esc(v):
    if v is None:
        return ""
    s = str(v).strip()
    if s in ("#N/A", "None"):
        return ""
    s = re.sub(r"\s+", " ", s)
    for a, b in ERRATAS:
        s = s.replace(a, b)
    for a, b in REPL:
        s = s.replace(a, b)
    return cursivas(superindices(s))


wb = openpyxl.load_workbook(XLSX, data_only=True)
out = []


def longtable(caption, colspec, headers, rows, label, nota=None):
    out.append(r"\begin{footnotesize}")
    out.append(r"\begin{longtable}{%s}" % colspec)
    out.append(r"\caption{%s\vspace{1ex}}\label{tab:%s}\\" % (caption, label))
    out.append(r"\toprule")
    out.append(" & ".join(r"\textbf{%s}" % h for h in headers) + r" \\")
    out.append(r"\midrule\endfirsthead")
    out.append(r"\multicolumn{%d}{@{}l@{}}{\textit{(continuación)}}\\" % len(headers))
    out.append(r"\toprule")
    out.append(" & ".join(r"\textbf{%s}" % h for h in headers) + r" \\")
    out.append(r"\midrule\endhead")
    out.append(r"\bottomrule\endfoot")
    for r in rows:
        out.append(" & ".join(r) + r" \\")
    out.append(r"\end{longtable}")
    if nota:
        out.append(r"\vspace{-1ex}\noindent\parbox{\textwidth}{\footnotesize %s}"
                   % nota)
    out.append(r"\end{footnotesize}")
    out.append("")


# Mitigaciones desde la hoja RISK (columnas 1-9 para UR, 11-19 para SR)
mit = {}
for row in wb["RISK"].iter_rows(min_row=2, values_only=True):
    for base in (0, 10):
        rid = row[base]
        if rid and str(rid).strip().startswith(("UR-", "SR-")):
            mit[str(rid).strip()] = row[base + 8]
mit.update(MIT_OVERRIDE)   # reasigna las mitigaciones corridas en el Excel

GUT = r"@{\hspace{4pt}}"
REQ_COLS = (r"|" + GUT + r"l >{\raggedright\arraybackslash}p{5.4cm} c c c "
            r">{\raggedright\arraybackslash}p{4.6cm}" + GUT + r"|")
REQ_HEAD = ["ID", "Requerimiento", "Prior.", "Riesgo", "Tipo", "Mitigación propuesta"]


def req_rows(sheet):
    rows = []
    for r in wb[sheet].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        rid = str(r[0]).strip()
        rows.append([esc(rid), esc(r[2]), esc(r[3]), esc(r[4]), esc(r[6]),
                     esc(mit.get(rid))])
    return rows


# ---------- Anexo A ----------------------------------------------------------
out.append(r"\subsection{Anexo A: requerimientos de usuario}")
out.append("")
longtable(
    "Requerimientos de usuario del proyecto BIFROST. Elaboración propia.",
    REQ_COLS, REQ_HEAD, req_rows("UR"), "anexoA",
    nota=(r"Nota. UR-13 se incorporó al cierre de la PDR, con posterioridad al "
          r"análisis de riesgos, por lo que todavía no tiene prioridad ni nivel "
          r"de riesgo asignados."))

# ---------- Anexo B ----------------------------------------------------------
out.append(r"\subsection{Anexo B: requerimientos de sistema}")
out.append("")
longtable(
    "Requerimientos de sistema del proyecto BIFROST. Elaboración propia.",
    REQ_COLS, REQ_HEAD, req_rows("SR"), "anexoB",
    nota=(r"Nota. La numeración salta de SR-10 a SR-13 porque SR-11 y SR-12 se "
          r"eliminaron al consolidar la línea base y sus identificadores no se "
          r"reutilizaron. SR-19 a SR-23 se incorporaron al cierre de la PDR y "
          r"todavía no tienen evaluación de riesgo. SR-07 conserva la redacción "
          r"aprobada en la SRR, cuando la línea base contemplaba el uso del "
          r"transponder en banda S de la plataforma; durante la PDR el enlace de "
          r"bajada se reasignó a una antena UHF propia de la carga útil, que es "
          r"la configuración que describe el cuerpo de este informe. La "
          r"actualización formal del requerimiento queda pendiente para la CDR."))

# ---------- Anexo C ----------------------------------------------------------
pairs = {}
for a, b, *_ in wb["TR"].iter_rows(min_row=2, values_only=True):
    if a and b and str(a).startswith("UR-") and str(b).startswith("SR-"):
        pairs.setdefault(str(a).strip(), set()).add(str(b).strip())

out.append(r"\subsection{Anexo C: matriz de trazabilidad de requerimientos}")
out.append("")
longtable(
    "Matriz de trazabilidad entre requerimientos de usuario y de sistema. "
    "Elaboración propia.",
    r"|" + GUT + r"l >{\raggedright\arraybackslash}p{12.2cm}" + GUT + r"|",
    ["Req. de usuario", "Requerimientos de sistema asociados"],
    [[esc(k), esc(", ".join(sorted(v)))] for k, v in sorted(pairs.items())],
    "anexoC",
    nota=(r"Nota. La matriz registra 35 relaciones. Los 21 requerimientos de "
          r"sistema aparecen al menos una vez, de modo que la cobertura es "
          r"completa en ambos sentidos."))

# ---------- Anexo D ----------------------------------------------------------
tests = [r for r in wb["TEST"].iter_rows(min_row=2, max_col=4, values_only=True) if r[0]]

out.append(r"\subsection{Anexo D: matriz preliminar de ensayos de verificación}")
out.append("")
longtable(
    "Matriz preliminar de ensayos de verificación. Elaboración propia.",
    r"|" + GUT + r"l >{\raggedright\arraybackslash}p{4.0cm} >{\raggedright\arraybackslash}p{8.6cm}" + GUT + r"|",
    ["ID", "Nombre del ensayo", "Descripción"],
    [[esc(r[0]), esc(r[2]), esc(r[3])] for r in tests],
    "anexoD",
    nota=(r"Nota. Los 19 ensayos corresponden a especificaciones de "
          r"verificación definidas durante la PDR y ninguno se ejecutó dentro "
          r"del período de la práctica. Las menciones a banda S se refieren a "
          r"los equipos de la plataforma, que actúan como fuente de "
          r"interferencia sobre la carga útil."))

with open(OUT, "w", encoding="utf-8") as f:
    f.write("% ARCHIVO GENERADO AUTOMÁTICAMENTE por content/gen_tablas.py\n")
    f.write("% No editar a mano: vuelve a correr el script si cambia el Excel.\n\n")
    f.write("\n".join(out))

print("ok ->", OUT, "| UR:", len(req_rows("UR")), "SR:", len(req_rows("SR")),
      "trazas:", sum(len(v) for v in pairs.values()), "ensayos:", len(tests))
