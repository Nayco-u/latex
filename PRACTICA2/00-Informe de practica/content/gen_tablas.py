"""Genera content/tablas_anexos.tex a partir de 07-TraceabilityRequirimentsMatrix.xlsx.

Vuelve a correrlo si el Excel cambia:  python3 content/gen_tablas.py
"""
import re
import openpyxl

XLSX = "/mnt/user-data/uploads/PRACTICA2/07-TraceabilityRequirimentsMatrix.xlsx"
OUT = "/root/informe/content/tablas_anexos.tex"

REPL = [
    ("\\", r"\textbackslash{}"), ("&", r"\&"), ("%", r"\%"), ("$", r"\$"),
    ("#", r"\#"), ("_", r"\_"), ("{", r"\{"), ("}", r"\}"),
    ("~", r"\textasciitilde{}"), ("^", r"\textasciicircum{}"),
    ("°", r"$^{\circ}$"), ("±", r"$\pm$"), ("≥", r"$\geq$"), ("≤", r"$\leq$"),
    ("<", r"$<$"), (">", r"$>$"), ("×", r"$\times$"),
    ("“", "``"), ("”", "''"), ("’", "'"),
]


def esc(v):
    if v is None:
        return ""
    s = str(v).strip()
    if s in ("#N/A", "None"):
        return ""
    for a, b in REPL:
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s)


wb = openpyxl.load_workbook(XLSX, data_only=True)
out = []


def longtable(caption, colspec, headers, rows, label):
    out.append(r"\begin{footnotesize}")
    out.append(r"\begin{longtable}{%s}" % colspec)
    out.append(r"\caption{%s\vspace{1ex}}\label{tab:%s}\\" % (caption, label))
    out.append(r"\toprule")
    out.append(" & ".join(r"\textbf{%s}" % h for h in headers) + r" \\")
    out.append(r"\midrule\endfirsthead")
    out.append(r"\multicolumn{%d}{@{}l@{}}{\textit{(continuación)}}\\" % len(headers))
    out.append(r"\toprule")
    out.append(" & ".join(r"\textbf{%s}" % h for h in headers) + r" \\")
    out.append(r"\midrule\endhead")
    out.append(r"\bottomrule\endfoot")
    for r in rows:
        out.append(" & ".join(r) + r" \\")
        out.append(r"\addlinespace[2pt]")
    out.append(r"\end{longtable}")
    out.append(r"\end{footnotesize}")
    out.append("")


# Mitigaciones desde la hoja RISK (columnas 1-9 para UR, 11-19 para SR)
mit = {}
for row in wb["RISK"].iter_rows(min_row=2, values_only=True):
    for base in (0, 10):
        rid = row[base]
        if rid and str(rid).strip().startswith(("UR-", "SR-")):
            mit[str(rid).strip()] = row[base + 8]

REQ_COLS = (r"@{}l >{\raggedright\arraybackslash}p{5.4cm} c c c "
            r">{\raggedright\arraybackslash}p{4.6cm}@{}")
REQ_HEAD = ["ID", "Requerimiento", "Prior.", "Riesgo", "Tipo", "Mitigación propuesta"]


def req_rows(sheet):
    rows = []
    for r in wb[sheet].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        rid = str(r[0]).strip()
        rows.append([esc(rid), esc(r[2]), esc(r[3]), esc(r[4]), esc(r[6]),
                     esc(mit.get(rid))])
    return rows


# ---------- Anexo A ----------------------------------------------------------
out.append(r"\subsection{Anexo A: requerimientos de usuario}")
out.append("")
longtable(
    "Requerimientos de usuario del proyecto BIFROST, con su prioridad, "
    "nivel de riesgo, tipo y mitigación definida. Elaboración propia a partir de "
    "la matriz de trazabilidad del proyecto.",
    REQ_COLS, REQ_HEAD, req_rows("UR"), "anexoA")

# ---------- Anexo B ----------------------------------------------------------
out.append(r"\subsection{Anexo B: requerimientos de sistema}")
out.append("")
longtable(
    "Requerimientos de sistema derivados de los requerimientos de "
    "usuario, con su prioridad, nivel de riesgo, tipo y mitigación definida. "
    "Elaboración propia a partir de la matriz de trazabilidad del proyecto.",
    REQ_COLS, REQ_HEAD, req_rows("SR"), "anexoB")

# ---------- Anexo C ----------------------------------------------------------
pairs = {}
for a, b, *_ in wb["TR"].iter_rows(min_row=2, values_only=True):
    if a and b and str(a).startswith("UR-") and str(b).startswith("SR-"):
        pairs.setdefault(str(a).strip(), set()).add(str(b).strip())

out.append(r"\subsection{Anexo C: matriz de trazabilidad de requerimientos}")
out.append("")
longtable(
    "Trazabilidad entre requerimientos de usuario y requerimientos de "
    "sistema. Elaboración propia.",
    r"@{}l >{\raggedright\arraybackslash}p{12.2cm}@{}",
    ["Req. de usuario", "Requerimientos de sistema asociados"],
    [[esc(k), esc(", ".join(sorted(v)))] for k, v in sorted(pairs.items())],
    "anexoC")

# ---------- Anexo D ----------------------------------------------------------
tests = [r for r in wb["TEST"].iter_rows(min_row=2, max_col=4, values_only=True) if r[0]]

out.append(r"\subsection{Anexo D: matriz preliminar de ensayos de verificación}")
out.append("")
longtable(
    "Ensayos de verificación especificados para las fases de diseño "
    "detallado y calificación del sistema. Elaboración propia.",
    r"@{}l >{\raggedright\arraybackslash}p{4.0cm} >{\raggedright\arraybackslash}p{8.6cm}@{}",
    ["ID", "Nombre del ensayo", "Descripción"],
    [[esc(r[0]), esc(r[2]), esc(r[3])] for r in tests],
    "anexoD")

with open(OUT, "w", encoding="utf-8") as f:
    f.write("% ARCHIVO GENERADO AUTOMÁTICAMENTE por content/gen_tablas.py\n")
    f.write("% No editar a mano: vuelve a correr el script si cambia el Excel.\n\n")
    f.write("\n".join(out))

print("ok ->", OUT, "| UR:", len(req_rows("UR")), "SR:", len(req_rows("SR")),
      "trazas:", sum(len(v) for v in pairs.values()), "ensayos:", len(tests))
